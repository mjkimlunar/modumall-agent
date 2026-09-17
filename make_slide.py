# -*- coding: utf-8 -*-
"""팀 공유용 발표 슬라이드 1장을 만든다.  `python make_slide.py`

슬라이드는 읽는 것이 아니라 훑는 것이므로 한 항목당 한 줄로 끊는다.
상세 근거는 아티팩트·노션 리포트가 담당한다.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).parent / "발표자료_모두몰에이전트.pptx"

FONT = "맑은 고딕"
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x89, 0x87, 0x81)
SURFACE = RGBColor(0xFC, 0xFC, 0xFB)
BAND = RGBColor(0xF0, 0xEF, 0xEC)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
AQUA = RGBColor(0x1B, 0xAF, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def para(tf, text, size=12, bold=False, color=INK2, space_after=6):
    first = not tf.paragraphs[0].runs and not tf.paragraphs[0].text
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = text
    p.space_after = Pt(space_after)
    for r in p.runs:
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    return p


def item(tf, head, tail, color=INK):
    """한 줄짜리 항목: 굵은 제목 + 얇은 한 줄 설명."""
    para(tf, head, size=13, bold=True, color=color, space_after=2)
    para(tf, tail, size=11, color=MUTED, space_after=14)


def column(slide, left, top, width, height, title, color):
    head = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.46))
    head.fill.solid(); head.fill.fore_color.rgb = color
    head.line.fill.background()
    tf = head.text_frame
    tf.margin_left, tf.margin_top, tf.margin_bottom = Inches(0.15), 0, 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, title, size=14, bold=True, color=WHITE, space_after=0)

    body = slide.shapes.add_textbox(Inches(left), Inches(top + 0.62),
                                    Inches(width), Inches(height - 0.62))
    b = body.text_frame
    b.word_wrap = True
    b.margin_left = Inches(0.06)
    return b


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = SURFACE

    # ── 제목 ────────────────────────────────────────────────
    t = slide.shapes.add_textbox(Inches(0.5), Inches(0.38), Inches(12.4), Inches(0.55)).text_frame
    para(t, "모두몰 고객응대 에이전트 — 성능 개선", size=26, bold=True, color=INK, space_after=0)

    s = slide.shapes.add_textbox(Inches(0.52), Inches(0.96), Inches(12.4), Inches(0.35)).text_frame
    para(s, "실패 사례를 하나씩 읽고 고침 · 판단 기준은 전부 매뉴얼에서 인용 · 평가 기준(evaluate.py)은 미수정",
         size=12, color=MUTED, space_after=0)

    W, H, TOP = 3.9, 5.5, 1.6

    # ── ① 주요 개선점 ───────────────────────────────────────
    b = column(slide, 0.5, TOP, W, H, "①  주요 개선점", BLUE)
    item(b, "라우팅 경계 규칙", "매뉴얼 조항을 그대로 인용해 예시화 (§1.1 · §2.4 · §5.4)")
    item(b, "조회 절차화", "세부 속성·재입고·반품 기간을 빠뜨리지 않게")
    item(b, "상품 검색 재설계", "부분 문자열 동점으로 엉뚱한 상품을 고르던 문제 해소")
    item(b, "오안내 3건 수정", "A/S 안내 · 반복 문의 이관 · 미확정을 확정으로 단정", color=ORANGE)
    para(b, "→ 셋 다 채점에 안 잡히던 버그", size=11, color=MUTED, space_after=0)

    # ── ② 지표 변화 ────────────────────────────────────────
    b = column(slide, 4.85, TOP, W, H, "②  지표 변화", AQUA)
    rows = [("지표", "기준선", "최종"),
            ("의도 분류 정확도", "0.942", "0.980"),
            ("의도 분류 macro F1", "0.945", "0.984"),
            ("1턴 답변 통과율", "57.3%", "77.1%"),
            ("실전 자동화율", "7.5%", "12.5%")]
    tbl = slide.shapes.add_table(len(rows), 3, Inches(4.85), Inches(TOP + 0.7),
                                 Inches(W), Inches(2.0)).table
    tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = \
        Inches(1.9), Inches(1.0), Inches(1.0)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(0.4)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = BAND if i == 0 else SURFACE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.07)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            for r in p.runs:
                r.font.name, r.font.size = FONT, Pt(11.5)
                r.font.bold = (i == 0) or (j == 2 and i > 0)
                r.font.color.rgb = INK if i == 0 else (AQUA if j == 2 else INK2)

    note = slide.shapes.add_textbox(Inches(4.85), Inches(TOP + 2.95),
                                    Inches(W), Inches(2.4)).text_frame
    note.word_wrap = True
    item(note, "라우팅 오분류 7건 → 2건", "경계 예시 3~4개만으로, 코드 수정 없이")
    item(note, "비용 +37%", "1회 호출 5,024 → 6,881 토큰 · 1만 건 ≈ 57,800원")
    item(note, "모든 수치는 3회 평균", "같은 입력도 ±9%p 흔들려 단일 측정은 쓰지 않음")

    # ── ③ 개선하지 못한 점 ──────────────────────────────────
    b = column(slide, 9.2, TOP, W, H, "③  개선하지 못한 점 / 이유", ORANGE)
    item(b, "남은 실패 6건 = 채점기 한계", "동의어·띄어쓰기·표기 차이, 도구 부르면 무조건 ANSWER 판정")
    item(b, "데이터 스키마 한계", "옵션별 재고·수령일이 조회에 없어 계산 불가")
    item(b, "자동화율 12.5%의 벽", "문의 77.5%가 상품명 자체를 말하지 않음")
    item(b, "멀티턴 2차 턴이 가장 약함", "1차 대비 약 30%p 낮음 · 한 번도 재지 않던 구간")
    para(b, "→ 재는 곳만 좋아지고, 안 재는 곳은 조용히 틀린 채로 남는다",
         size=11, bold=True, color=INK2, space_after=0)

    prs.save(OUT)
    print("생성:", OUT)


if __name__ == "__main__":
    main()
